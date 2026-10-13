from pathlib import Path

import pandas as pd
import pytest
from docx import Document as new_docx

from app.core.exceptions import UnsupportedFileTypeError
from app.infrastructure.loaders.docx_loader import DocxLoader
from app.infrastructure.loaders.factory import LoaderFactory
from app.infrastructure.loaders.spreadsheet_loader import SpreadsheetLoader
from app.infrastructure.loaders.text_loader import TextLoader
from app.infrastructure.ocr.scan_detector import HeuristicScanDetector


class FakeOcr:
    def image_to_text(self, image):
        return "texto reconhecido"


def test_spreadsheet_rows_become_key_value_blocks(tmp_path: Path):
    path = tmp_path / "disciplinas.xlsx"
    frame = pd.DataFrame(
        {"Código": [f"GEP{i:03d}" for i in range(50)], "Disciplina": [f"Disciplina {i}" for i in range(50)]}
    )
    frame.to_excel(path, index=False, sheet_name="Grade")
    loaded = SpreadsheetLoader(block_chars=300).load(path)

    assert len(loaded.sections) > 1
    first = loaded.sections[0]
    assert first.sheet == "Grade" and first.kind == "table"
    assert "Código: GEP000; Disciplina: Disciplina 0" in first.text
    assert first.heading.startswith("Planilha Grade (linhas 2-")


def test_csv_with_latin1_encoding(tmp_path: Path):
    path = tmp_path / "notas.csv"
    path.write_bytes("Aluno;Situação\nJoão;Aprovado\nMaria;Reprovação\n".encode("latin-1"))
    loaded = SpreadsheetLoader().load(path)
    assert "Situação: Aprovado" in loaded.sections[0].text


def test_docx_groups_by_heading_and_reads_tables(tmp_path: Path):
    path = tmp_path / "regulamento.docx"
    document = new_docx()
    document.add_heading("Capítulo 1 - Matrícula", level=1)
    document.add_paragraph("A matrícula é semestral.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Prazo", "Março"
    table.cell(1, 0).text, table.cell(1, 1).text = "Local", "Secretaria"
    document.add_heading("Capítulo 2 - Trancamento", level=1)
    document.add_paragraph("O trancamento exige requerimento.")
    document.save(path)

    sections = DocxLoader().load(path).sections
    assert [s.heading for s in sections] == ["Capítulo 1 - Matrícula", "Capítulo 2 - Trancamento"]
    assert "Prazo | Março" in sections[0].text


def test_markdown_sections(tmp_path: Path):
    path = tmp_path / "faq.md"
    path.write_text("# Bolsas\nInformações.\n\n## Auxílio\nDetalhes.", encoding="utf-8")
    sections = TextLoader().load(path).sections
    assert [s.heading for s in sections] == ["Bolsas", "Auxílio"]


def test_factory_resolves_by_extension():
    factory = LoaderFactory.default(FakeOcr(), HeuristicScanDetector(), 150, "soffice")
    assert factory.supports("Edital.PDF")
    assert ".docx" in factory.supported_extensions and ".csv" in factory.supported_extensions
    with pytest.raises(UnsupportedFileTypeError):
        factory.for_path(Path("programa.exe"))
