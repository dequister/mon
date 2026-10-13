from pathlib import Path

from pptx import Presentation

from app.domain.entities import DocumentSection, LoadedDocument
from app.infrastructure.loaders.base import BaseLoader


class PptxLoader(BaseLoader):
    """Uma seção por slide, incluindo tabelas e anotações do apresentador."""

    extensions = (".pptx",)

    def _extract(self, path: Path) -> LoadedDocument:
        presentation = Presentation(str(path))
        sections: list[DocumentSection] = []
        for number, slide in enumerate(presentation.slides, start=1):
            lines: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines.append(shape.text_frame.text)
                elif getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        lines.append(" | ".join(cell.text.strip() for cell in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"Notas: {notes}")
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else None
            sections.append(DocumentSection(text="\n".join(lines), page=number, heading=title or None))
        return LoadedDocument(sections=sections, page_count=len(presentation.slides))
